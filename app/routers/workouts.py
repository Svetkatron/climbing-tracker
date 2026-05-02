from fastapi import APIRouter, Depends, HTTPException, status
from sqlalchemy.orm import Session
from typing import List
from datetime import datetime
from app.database import get_db
from app import models, schemas, auth
from datetime import datetime, timedelta
from calendar import monthrange
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import io

router = APIRouter(prefix="/workouts", tags=["workouts"])

@router.get("/export/excel")
def export_workouts_excel(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_active_user)
):
    """Экспорт тренировок в Excel файл"""
    
    # Получаем все тренировки пользователя
    workouts = db.query(models.Workout).filter(
        models.Workout.user_id == current_user.id
    ).order_by(models.Workout.date.desc()).all()
    
    # Создаём Excel файл
    wb = Workbook()
    ws = wb.active
    ws.title = "Мои тренировки"
    
    # Заголовки
    headers = ["Дата", "Трасса", "Сложность", "Попытки", "Успех", "Заметки", "Соревнование"]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2c3e50", end_color="2c3e50", fill_type="solid")
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    
    # Заполняем данными
    for row, workout in enumerate(workouts, 2):
        ws.cell(row=row, column=1, value=workout.date.strftime("%d.%m.%Y %H:%M") if workout.date else "")
        ws.cell(row=row, column=2, value=workout.route.name if workout.route else "Неизвестно")
        ws.cell(row=row, column=3, value=workout.route.grade if workout.route else "")
        ws.cell(row=row, column=4, value=workout.attempts)
        ws.cell(row=row, column=5, value="✅ Да" if workout.success else "❌ Нет")
        ws.cell(row=row, column=6, value=workout.notes or "")
        ws.cell(row=row, column=7, value="🏆 Да" if workout.is_competition else "❌ Нет")
    
    # Настраиваем ширину колонок
    column_widths = [20, 25, 12, 10, 10, 35, 12]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    
    # Сохраняем в буфер
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    # Возвращаем файл
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=trainings_{current_user.username}.xlsx"}
    )

@router.post("/", response_model=schemas.WorkoutResponse, status_code=status.HTTP_201_CREATED)
def create_workout(
    workout: schemas.WorkoutCreate,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_active_user)
):
    route = db.query(models.Route).filter(models.Route.id == workout.route_id).first()
    if not route:
        raise HTTPException(status_code=404, detail="Route not found")
    
    db_workout = models.Workout(
        user_id=current_user.id,
        route_id=workout.route_id,
        attempts=workout.attempts,
        success=workout.success,
        notes=workout.notes,
        date=workout.date or datetime.utcnow()
    )
    
    db.add(db_workout)
    db.commit()
    db.refresh(db_workout)
    return db_workout

@router.put("/{workout_id}", response_model=schemas.WorkoutResponse)
def update_workout(
    workout_id: int,
    workout_update: schemas.WorkoutUpdate,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_active_user)
):
    workout = db.query(models.Workout).filter(
        models.Workout.id == workout_id,
        models.Workout.user_id == current_user.id
    ).first()
    if not workout:
        raise HTTPException(status_code=404, detail="Workout not found")
    
    # Если изменилась дата, сбрасываем reminder_sent
    if workout_update.date and workout_update.date != workout.date:
        workout_update.reminder_sent = False
    
    update_data = workout_update.model_dump(exclude_unset=True)
    for key, value in update_data.items():
        setattr(workout, key, value)
    
    db.commit()
    db.refresh(workout)
    return workout

@router.delete("/{workout_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_workout(
    workout_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_active_user)
):
    workout = db.query(models.Workout).filter(
        models.Workout.id == workout_id,
        models.Workout.user_id == current_user.id
    ).first()
    if not workout:
        raise HTTPException(status_code=404, detail="Workout not found")
    
    db.delete(workout)
    db.commit()
    return None

@router.get("/calendar/{year}/{month}")
def get_calendar(
    year: int,
    month: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_active_user)
):
    """
    Тут получаем календарь тренировок за месяц
    Возвращает дни месяца и количество тренировок в каждый день
    """
    # Проверяем корректность месяца
    if month < 1 or month > 12:
        raise HTTPException(status_code=400, detail="Month must be between 1 and 12")
    
    # Получаем первый и последний день месяца
    start_date = datetime(year, month, 1)
    last_day = monthrange(year, month)[1]
    end_date = datetime(year, month, last_day, 23, 59, 59)
    
    # Получаем тренировки за месяц
    workouts = db.query(models.Workout).filter(
        models.Workout.user_id == current_user.id,
        models.Workout.date >= start_date,
        models.Workout.date <= end_date
    ).all()
    
    # Группируем тренировки по дням
    calendar_data = {}
    for day in range(1, last_day + 1):
        calendar_data[day] = {
            "date": f"{year}-{month:02d}-{day:02d}",
            "workouts_count": 0,
            "total_attempts": 0,
            "successful_attempts": 0,
            "workouts": []
        }
    
    for workout in workouts:
        day = workout.date.day
        calendar_data[day]["workouts_count"] += 1
        calendar_data[day]["total_attempts"] += workout.attempts
        if workout.success:
            calendar_data[day]["successful_attempts"] += workout.attempts
        calendar_data[day]["workouts"].append({
            "id": workout.id,
            "route_name": workout.route.name if workout.route else "Unknown",
            "grade": workout.route.grade if workout.route else "Unknown",
            "attempts": workout.attempts,
            "success": workout.success,
            "notes": workout.notes,
            "is_competition": workout.is_competition
        })
    
    # Формируем результат
    result = []
    for day in range(1, last_day + 1):
        result.append({
            "day": day,
            "date": calendar_data[day]["date"],
            "workouts_count": calendar_data[day]["workouts_count"],
            "total_attempts": calendar_data[day]["total_attempts"],
            "successful_attempts": calendar_data[day]["successful_attempts"],
            "success_rate": round(
                calendar_data[day]["successful_attempts"] / calendar_data[day]["total_attempts"] * 100, 1
            ) if calendar_data[day]["total_attempts"] > 0 else 0,
            "workouts": calendar_data[day]["workouts"]
        })
    
    return {
        "year": year,
        "month": month,
        "days_in_month": last_day,
        "total_workouts": len(workouts),
        "calendar": result
    }


@router.get("/calendar/stats")
def get_calendar_stats(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_active_user)
):
    """
    Получить общую статистику активности по месяцам
    """
    # Получаем все тренировки пользователя
    workouts = db.query(models.Workout).filter(
        models.Workout.user_id == current_user.id
    ).all()
    
    if not workouts:
        return {
            "total_workouts": 0,
            "total_days": 0,
            "streak_current": 0,
            "streak_best": 0,
            "months": []
        }
    
    # Группируем по месяцам
    months_stats = {}
    workout_dates = []
    
    for workout in workouts:
        year = workout.date.year
        month = workout.date.month
        day = workout.date.day
        date_key = f"{year}-{month:02d}-{day:02d}"
        
        if date_key not in workout_dates:
            workout_dates.append(date_key)
        
        month_key = f"{year}-{month:02d}"
        if month_key not in months_stats:
            months_stats[month_key] = {
                "year": year,
                "month": month,
                "workouts_count": 0,
                "days_with_workouts": set()
            }
        months_stats[month_key]["workouts_count"] += 1
        months_stats[month_key]["days_with_workouts"].add(day)
    
    # Сортируем даты для расчёта серий
    workout_dates.sort()
    
    # Расчёт текущей и максимальной серии (стрейка)
    current_streak = 0
    best_streak = 0
    temp_streak = 0
    prev_date = None
    
    for date_str in workout_dates:
        if prev_date is None:
            temp_streak = 1
        else:
            prev = datetime.strptime(prev_date, "%Y-%m-%d")
            curr = datetime.strptime(date_str, "%Y-%m-%d")
            diff = (curr - prev).days
            if diff == 1:
                temp_streak += 1
            else:
                if temp_streak > best_streak:
                    best_streak = temp_streak
                temp_streak = 1
        prev_date = date_str
    
    if temp_streak > best_streak:
        best_streak = temp_streak
    
    # Проверяем текущую серию (до сегодня)
    last_workout = datetime.strptime(workout_dates[-1], "%Y-%m-%d")
    today = datetime.utcnow()
    days_since_last = (today - last_workout).days
    
    if days_since_last == 0:
        # Сегодня была тренировка, считаем серию
        current_streak = 1
        # Идём назад
        check_date = last_workout - timedelta(days=1)
        while check_date.strftime("%Y-%m-%d") in workout_dates:
            current_streak += 1
            check_date -= timedelta(days=1)
    else:
        current_streak = 0
    
    # Формируем результат по месяцам
    months_result = []
    for month_key in sorted(months_stats.keys(), reverse=True):
        stats = months_stats[month_key]
        months_result.append({
            "year": stats["year"],
            "month": stats["month"],
            "workouts_count": stats["workouts_count"],
            "active_days": len(stats["days_with_workouts"])
        })
    
    return {
        "total_workouts": len(workouts),
        "total_days": len(workout_dates),
        "streak_current": current_streak,
        "streak_best": best_streak,
        "months": months_result
    }